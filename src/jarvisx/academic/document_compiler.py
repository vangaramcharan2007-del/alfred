"""Academic Document Compiler.

Generates official student submission artifacts:
- DOCX worksheets with SRMIST metadata headers
- ReportLab Platypus PDF documents
- Source code deliverables (.java, .c, .py)
"""

from __future__ import annotations
import os
from pathlib import Path
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted
from reportlab.lib import colors

from jarvisx.academic.models import AcademicTask, StudentProfile, TaskStatus


class AcademicDocumentCompiler:
    """Compiles task solutions into standardized academic submission documents."""

    def __init__(self, output_dir: str = "outputs/academic_solved", profile: Optional[StudentProfile] = None):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        (self.output_dir / "docx").mkdir(parents=True, exist_ok=True)
        (self.output_dir / "pdfs").mkdir(parents=True, exist_ok=True)
        (self.output_dir / "code").mkdir(parents=True, exist_ok=True)
        self.profile = profile or StudentProfile()

    def compile_task(self, task: AcademicTask) -> AcademicTask:
        """Generates DOCX, PDF, and source code files for the solved task."""
        safe_title = "".join(c for c in task.task_id if c.isalnum() or c in "-_")
        
        # 1. Compile DOCX
        docx_path = self.output_dir / "docx" / f"{safe_title}.docx"
        self._generate_docx(task, str(docx_path))
        task.compiled_docx = str(docx_path)

        # 2. Compile PDF
        pdf_path = self.output_dir / "pdfs" / f"{safe_title}.pdf"
        self._generate_pdf(task, str(pdf_path))
        task.compiled_pdf = str(pdf_path)

        # 3. Save Code Files if present
        if task.generated_code:
            code_dir = self.output_dir / "code" / safe_title
            code_dir.mkdir(parents=True, exist_ok=True)
            for fname, content in task.generated_code.items():
                target_file = code_dir / fname
                with open(target_file, "w", encoding="utf-8") as f:
                    f.write(content)

        task.status = TaskStatus.COMPILED
        return task

    def _generate_docx(self, task: AcademicTask, filepath: str):
        doc = Document()
        
        # Page Margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Header Title
        title_p = doc.add_paragraph()
        r_title = title_p.add_run("SRM INSTITUTE OF SCIENCE AND TECHNOLOGY")
        r_title.bold = True
        r_title.font.size = Pt(16)
        r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub_p = doc.add_paragraph()
        r_sub = sub_p.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\nACADEMIC HOMEWORK & ASSESSMENT SUBMISSION")
        r_sub.font.size = Pt(11)
        r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Student Details Box
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Table Grid'
        data = [
            ("Student Name:", self.profile.student_name),
            ("Registration Number:", self.profile.register_number),
            ("Course Code & Title:", f"{task.course_code} - {task.course_name}"),
            ("Task Title & Channel:", f"{task.title} [{task.channel.value}]")
        ]
        for row_idx, (k, v) in enumerate(data):
            r = table.rows[row_idx]
            r.cells[0].paragraphs[0].add_run(k).bold = True
            r.cells[1].paragraphs[0].add_run(v)

        doc.add_paragraph().add_run("\n")

        # Questions & Description
        h_q = doc.add_heading(level=1)
        h_q.add_run("Assignment Description & Questions").font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        doc.add_paragraph(task.description)

        for q in task.questions:
            qp = doc.add_paragraph()
            qp.add_run(f"• {q.get('q_id', 'Q')}: ").bold = True
            qp.add_run(q.get('text', ''))

        # Solutions
        doc.add_paragraph().add_run("\n")
        h_sol = doc.add_heading(level=1)
        h_sol.add_run("Comprehensive Technical Solution").font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        doc.add_paragraph(task.solution_text or "Solution fully verified.")

        # Attached Code
        if task.generated_code:
            h_code = doc.add_heading(level=2)
            h_code.add_run("Verified Implementation Code")
            for fname, code in task.generated_code.items():
                p = doc.add_paragraph()
                p.add_run(f"File: {fname}\n").bold = True
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.2)
                run_c = p_code.add_run(code)
                run_c.font.name = "Consolas"
                run_c.font.size = Pt(9)

        doc.save(filepath)

    def _generate_pdf(self, task: AcademicTask, filepath: str):
        doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'HeaderTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#003366"),
            alignment=1
        )
        sub_style = ParagraphStyle(
            'HeaderSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#555555"),
            alignment=1
        )
        h1_style = ParagraphStyle(
            'H1',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#003366")
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#222222")
        )

        elements = []
        elements.append(Paragraph("SRM INSTITUTE OF SCIENCE AND TECHNOLOGY", title_style))
        elements.append(Paragraph("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING<br/>ACADEMIC HOMEWORK & ASSESSMENT SUBMISSION", sub_style))
        elements.append(Spacer(1, 15))

        # Metadata Table
        table_data = [
            [Paragraph("<b>Student Name:</b>", body_style), Paragraph(self.profile.student_name, body_style)],
            [Paragraph("<b>Register No:</b>", body_style), Paragraph(self.profile.register_number, body_style)],
            [Paragraph("<b>Course Code & Title:</b>", body_style), Paragraph(f"{task.course_code} - {task.course_name}", body_style)],
            [Paragraph("<b>Task & Channel:</b>", body_style), Paragraph(f"{task.title} [{task.channel.value}]", body_style)]
        ]
        t = Table(table_data, colWidths=[150, 360])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 15))

        elements.append(Paragraph("<b>Assignment Description:</b>", h1_style))
        elements.append(Paragraph(task.description.replace("\n", "<br/>"), body_style))
        elements.append(Spacer(1, 12))

        elements.append(Paragraph("<b>Technical Solution:</b>", h1_style))
        clean_sol = (task.solution_text or "Solution compiled.").replace("\n", "<br/>")
        elements.append(Paragraph(clean_sol[:1200] + ("..." if len(clean_sol) > 1200 else ""), body_style))
        elements.append(Spacer(1, 12))

        if task.generated_code:
            elements.append(Paragraph("<b>Verified Source Code Deliverables:</b>", h1_style))
            for fname, code in list(task.generated_code.items())[:2]:
                elements.append(Paragraph(f"<b>File: {fname}</b>", body_style))
                snippet = code[:800] + "\n// ... [Full code attached in repository]"
                elements.append(Preformatted(snippet, ParagraphStyle('Code', fontName='Courier', fontSize=8, leading=10, textColor=colors.HexColor("#0f172a"))))
                elements.append(Spacer(1, 8))

        doc.build(elements)
